"""Standardowy profil zużycia G11 operatora Enea Operator (załącznik nr 4 do IRiESD).

Plik XLSX na każdy rok. Układ arkusza nie jest nigdzie opisany w formie maszynowej, więc parser
próbuje kolejnych interpretacji i przyjmuje pierwszą, która przechodzi kontrolę jakości:
  * układ „długi”: wiersz nagłówka z kolumną G11 (także nagłówek dwuwierszowy); doba z kolumny daty
    (data, tekst albo liczba seryjna Excela) lub z kolumn Rok/Miesiąc/Dzień; okres z kolumny
    godziny/kwadransa (numer 1–25 / 1–100, „2A”, „00:00”, „00:00-01:00”, czas Excela), ze znacznika
    czasu w kolumnie daty albo z kolejności wierszy,
  * układ „szeroki”: wiersze = doby, kolumny = godziny (1–25) albo kwadranse (1–100), w arkuszu G11,
  * arkusz na miesiąc – poprawne doby z kolejnych arkuszy są łączone.
Etykiety końca okresu („01:00”…„24:00”, znaczniki końca godziny) są rozpoznawane i przesuwane
na początek okresu. Wiersze sum („Suma”, „Razem”) są pomijane.

Kontrola każdej doby: liczba okresów zgodna z długością doby (23/24/25 h albo 92/96/100 kwadransów),
energia w co najmniej 20 różnych godzinach, brak wartości ujemnych. Do 5% wadliwych lub brakujących
dób uzupełniamy średnią (miesiąc, typ dnia, slot) z poprawnych dób; przy większej liczbie skrypt
zatrzymuje się z diagnozą układu pliku.

Wynik: data/profil/enea_g11_RRRR.csv z kolumnami data, slot (1–96 wg czasu zegarowego), wartosc
(udział w roku w ppm – suma roku = 1 000 000), zrodlo. W dobie jesiennej zmiany czasu sloty
02:00–02:45 zawierają energię obu godzin 2:00–3:00. Normalizację do sumy 1 w roku robi Power Query.
"""
from __future__ import annotations

import re
import statistics
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path

import requests

from kalendarz import typ_dnia
from wspolne import (KATALOG_PROFIL, KATALOG_ZRODLA, TZ_PL, UTC, BladPobierania, czytaj_csv, fmt_liczba,
                     kwadranse_doby, log, nowa_sesja, poczatek_doby_utc, pobierz, zakres_dni, zapisz_csv_atomowo,
                     zapisz_stan)

# Adresy z https://www.operator.enea.pl/uslugi-dystrybucyjne/iriesd („Pliki do pobrania”).
# Na kolejny rok dopisz nowy adres – identyfikatory /media/... nie są przewidywalne.
PROFILE_URL = {
    2024: "https://www.operator.enea.pl/media/557/zalacznik4iriesdprofile2024xlsx.xlsx",
    2025: "https://www.operator.enea.pl/media/558/zalacznik4iriesdprofile2025xlsx.xlsx",
    2026: "https://www.operator.enea.pl/media/3253/2025/zalacznik4iriesdprofile2026xlsx.xlsx",
}
GRUPA = "G11"
NAGLOWEK = ["data", "slot", "wartosc", "zrodlo"]
WERSJA_PARSERA = 4           # zmiana wersji wymusza ponowne przetworzenie zapisanych plików XLSX
OPISY: dict[int, str] = {}   # rok -> rozpoznany układ pliku (do logu i data/aktualizacja.csv)
MAKS_WADLIWYCH = 0.05        # odsetek wadliwych dób, powyżej którego układ uznajemy za nierozpoznany
SUMA_ROKU = 1_000_000.0      # wartości w ppm rocznego zużycia

RE_GRUPA = re.compile(r"(?<![0-9A-Za-z])G\s?-?\s?11(?![0-9A-Za-z])", re.I)
RE_POMIN = re.compile(r"\b(suma|razem|ogółem|total)\b", re.I)
RE_NAGL_DATA = re.compile(r"\b(data|doba|date)\b", re.I)
RE_NAGL_OKRES = re.compile(r"godz|hour|kwadr|okres|czas|^\s*h\s*$|^\s*nr\b", re.I)
RE_ROK = re.compile(r"^\s*(rok|year)\b", re.I)
RE_MIES = re.compile(r"^\s*(miesi|month)", re.I)
RE_DZIEN = re.compile(r"^\s*(dzie[nń]|day)\s*$", re.I)  # „Dzień”, ale nie „Dzień tygodnia”
_RE_CZAS = re.compile(r"^(\d{1,2}):(\d{2})(?::\d{2})?(?:\s*[-–]\s*\d{1,2}:\d{2}(?::\d{2})?)?\s*([ab*])?$", re.I)
_RE_ZAKRES_GODZ = re.compile(r"^(\d{1,2})\s*[-–]\s*(\d{1,2})$")
_RE_NUMER = re.compile(r"^(?:h|godz\.?|godzina|kw\.?|kwadrans)?\s*(\d{1,3})\s*([ab*])?$", re.I)
EXCEL_EPOKA = date(1899, 12, 30)


class BladProfilu(RuntimeError):
    pass


@dataclass
class Wpis:
    minuty: int | None        # początek okresu (minuty od północy, czas zegarowy) – tryb „min”
    klucz: tuple | None       # porządek okresu w dobie – tryby „nr” i „kolejnosc”
    wartosc: float | None     # None = pusta komórka


def plik_surowy(rok: int) -> Path:
    return KATALOG_ZRODLA / f"enea_profil_{rok}.xlsx"


def plik_profilu(rok: int) -> Path:
    return KATALOG_PROFIL / f"enea_g11_{rok}.csv"


# ------------------------------------------------------------------ komórki

def _tekst(v) -> str:
    return "" if v is None else str(v).strip()


def _jako_date(v, serial: bool = False) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str) and v.strip():
        s = v.strip().split()[0]
        for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d-%m-%Y", "%Y.%m.%d", "%d/%m/%Y", "%Y/%m/%d"):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                pass
    if serial and isinstance(v, (int, float)) and not isinstance(v, bool) and 20000 <= v <= 80000:
        return EXCEL_EPOKA + timedelta(days=int(v))
    return None


def _jako_znacznik(v) -> datetime | None:
    """Data z godziną (datetime albo tekst „RRRR-MM-DD GG:MM”, także „24:00”)."""
    if isinstance(v, datetime):
        return v
    if isinstance(v, str):
        s = " ".join(v.split())
        plus = timedelta(0)
        if re.search(r"\s24:00(:00)?$", s):
            s, plus = re.sub(r"\s24:00(:00)?$", " 00:00", s), timedelta(days=1)
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M"):
            try:
                return datetime.strptime(s, fmt) + plus
            except ValueError:
                pass
    return None


def _jako_liczba(v) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace(",", ".").replace("\xa0", "").replace(" ", "")
        try:
            return float(s)
        except ValueError:
            return None
    return None


def _pod(znak) -> int:
    """Podokres powtórzonej godziny w dobie zmiany czasu: 2A → 1, 2B → 2."""
    return {"a": 1, "*": 1, "b": 2}.get((znak or "").lower(), 0)


def _okres(v) -> tuple[str, int, int] | None:
    """Okres z komórki: ('min', minuty od północy, podokres) albo ('nr', numer, podokres)."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (time, datetime)):
        return "min", v.hour * 60 + v.minute, 0
    if isinstance(v, (int, float)):
        if float(v).is_integer() and 0 <= v <= 100:
            return "nr", int(v), 0
        if 0 < v < 1 and abs(v * 96 - round(v * 96)) < 1e-6:  # czas Excela (ułamek doby, pełne kwadranse)
            return "min", round(v * 1440), 0
        return None
    s = str(v).strip()
    m = _RE_CZAS.match(s)
    if m:
        return "min", int(m.group(1)) * 60 + int(m.group(2)), _pod(m.group(3))
    m = _RE_ZAKRES_GODZ.match(s)
    if m and (int(m.group(2)) - int(m.group(1))) % 24 == 1:
        return "min", int(m.group(1)) * 60, 0
    m = _RE_NUMER.match(s)
    if m and int(m.group(1)) <= 100:
        return "nr", int(m.group(1)), _pod(m.group(2))
    return None


def _okres_w_trybie(v, rodzaj: str) -> tuple[str, int, int] | None:
    o = _okres(v)
    if o is None:
        return None
    if rodzaj == "min" and o[0] == "nr" and isinstance(v, (int, float)) and 0 <= v <= 1:
        return "min", round(v * 1440), 0  # 0 i 1 w kolumnie czasu Excela = 00:00 i 24:00
    return o if o[0] == rodzaj else None


def _litera(j: int) -> str:
    s = ""
    j += 1
    while j:
        j, r = divmod(j - 1, 26)
        s = chr(65 + r) + s
    return s


def _podglad(wiersze: list[tuple], ile: int = 3) -> str:
    niepuste = [w for w in wiersze if any(c is not None for c in w)][:ile]
    return " | ".join(", ".join(repr(c)[:22] for c in w[:8]) for w in niepuste)


def _dominujacy_krok(roznice: list[int]) -> int | None:
    dodatnie = [r for r in roznice if r > 0]
    return Counter(dodatnie).most_common(1)[0][0] if dodatnie else None


def _poczatek_z_konca(koniec_lok: datetime, wystapienie: int, krok: int) -> datetime:
    """Lokalny (naiwny) koniec okresu -> lokalny początek; różnica liczona w UTC, więc doby zmiany czasu
    wychodzą poprawnie. Drugie wystąpienie tej samej godziny w dobie jesiennej to czas zimowy (fold=1)."""
    utc = koniec_lok.replace(tzinfo=TZ_PL, fold=min(wystapienie, 1)).astimezone(UTC)
    return (utc - timedelta(minutes=krok)).astimezone(TZ_PL).replace(tzinfo=None)


# ------------------------------------------------------------------ składanie doby

def _sloty_zegarowe(dzien: date, liczba: int, krok_min: int) -> list[int] | None:
    """Kolejne okresy doby (godziny lub kwadranse) -> sloty zegarowe 1–96 (z powtórzeniem w dzień jesiennej zmiany).

    Zwraca None, gdy liczba okresów nie pasuje do długości doby.
    """
    n_kw = kwadranse_doby(dzien)
    na_okres = krok_min // 15
    oczekiwane = n_kw // na_okres
    start = poczatek_doby_utc(dzien)
    if liczba == oczekiwane:
        sloty = []
        for k in range(liczba):
            lok = (start + timedelta(minutes=krok_min * k)).astimezone(TZ_PL)
            sloty.append(lok.hour * 4 + lok.minute // 15 + 1)
        return sloty
    if liczba == 96 // na_okres:  # doba zmiany czasu zapisana jak zwykła (np. pusta godzina 2:00–3:00)
        return [k * na_okres + 1 for k in range(liczba)]
    return None


def _sloty_doby(dzien: date, wpisy: list[Wpis], tryb: str, krok: int) -> tuple[dict[int, float] | None, str | None]:
    """Wartości doby -> {slot zegarowy: energia}; przy niezgodności (None, powód)."""
    n_kw = kwadranse_doby(dzien)
    na_okres = krok // 15
    suma: dict[int, float] = defaultdict(float)
    if tryb == "min":
        for m, ile in Counter(w.minuty for w in wpisy).items():
            if m is None or m % krok or not 0 <= m < 1440:
                return None, f"nieprawidłowy początek okresu ({m} min)"
            if ile > 1 and not (n_kw == 100 and 120 <= m < 180):
                return None, f"powtórzony okres {m // 60:02d}:{m % 60:02d}"
        godziny = len({w.minuty // 60 for w in wpisy})
        if godziny < (23 if n_kw == 92 else 24):
            return None, f"okresy tylko w {godziny} godzinach"
        for w in wpisy:
            if w.wartosc is not None:
                for j in range(na_okres):
                    suma[w.minuty // 15 + 1 + j] += w.wartosc / na_okres
    else:
        wpisy = sorted(wpisy, key=lambda w: w.klucz) if tryb == "nr" else list(wpisy)
        wartosci = [w.wartosc for w in wpisy]
        oczekiwane = n_kw // na_okres
        while len(wartosci) > oczekiwane and None in wartosci:  # puste kolumny ponad długość doby (np. 25. godzina)
            del wartosci[len(wartosci) - 1 - wartosci[::-1].index(None)]
        if len(wartosci) > oczekiwane and n_kw == 92:
            # doba wiosenna zapisana jak zwykła, z zerem w miejscu nieistniejącej godziny – tak robi Enea
            # (godziny 1–24 numerowane wg końca okresu: kolumna „2” = 0, a 1:00–2:00 CET jest w kolumnie „3”);
            # usuwamy to zero i składamy 23 godziny kolejno w UTC
            nadmiar = len(wartosci) - oczekiwane
            zera = [i for i in range(min(len(wartosci), 16 // na_okres)) if wartosci[i] == 0][:nadmiar]
            if len(zera) == nadmiar:
                for i in reversed(zera):
                    del wartosci[i]
        sloty = _sloty_zegarowe(dzien, len(wartosci), krok)
        if sloty is None:
            return None, f"{len(wartosci)} okresów, a doba ma {oczekiwane}"
        for s0, v in zip(sloty, wartosci):
            if v is not None:
                for j in range(na_okres):
                    suma[s0 + j] += v / na_okres
    if any(v < 0 for v in suma.values()):
        return None, "ujemne wartości"
    z_energia = {(s - 1) // 4 for s, v in suma.items() if v > 0}
    if len(z_energia) < 20:
        return None, f"energia tylko w {len(z_energia)} godzinach"
    return dict(suma), None


def _ocen(doby: dict[date, list[Wpis]], tryb: str, krok: int | None):
    if krok not in (15, 60):
        return f"nie ustalono długości okresu (krok {krok})"
    dobre, zle = {}, {}
    for d in sorted(doby):
        sloty, powod = _sloty_doby(d, doby[d], tryb, krok)
        if sloty is None:
            zle[d] = powod
        else:
            dobre[d] = sloty
    return dobre, zle


def _krok_z_liczby(doby: dict[date, list[Wpis]]) -> int | None:
    if not doby:
        return None
    mediana = statistics.median(len(w) for w in doby.values())
    return 15 if mediana >= 88 else 60 if 20 <= mediana <= 27 else None


# ------------------------------------------------------------------ układ długi

def _kolumna_daty(probka: list[tuple], nagl: list[str], kolumny: list[int]) -> tuple[int | None, bool]:
    oceny = []
    for j in kolumny:
        serial = j < len(nagl) and bool(RE_NAGL_DATA.search(nagl[j]))
        ile = sum(1 for r in probka if j < len(r) and _jako_date(r[j], serial))
        oceny.append((ile, serial, -j, j))
    oceny.sort(reverse=True)
    if oceny and oceny[0][0] >= 3:
        return oceny[0][3], oceny[0][1]
    return None, False


def _parsuj_dlugi(dane: list[tuple], rok: int, nagl: list[str], kg: int):
    probka = [r for r in dane[:400] if any(c is not None for c in r)]
    if len(probka) < 20:
        return "za mało wierszy danych"
    szer = max(len(r) for r in probka)
    nagl = nagl + [""] * (szer - len(nagl))
    kolumny = [j for j in range(szer) if j != kg]

    # --- doba: kolumna daty albo Rok/Miesiąc/Dzień
    kd, serial = _kolumna_daty(probka, nagl, kolumny)
    podzial = None
    if kd is None:
        kr = next((j for j in kolumny if RE_ROK.search(nagl[j])), None)
        km = next((j for j in kolumny if RE_MIES.search(nagl[j])), None)
        kdz = next((j for j in kolumny if RE_DZIEN.search(nagl[j])), None)
        if None in (kr, km, kdz):
            return "brak kolumny daty"
        podzial = (kr, km, kdz)

    def doba(r) -> date | None:
        if kd is not None:
            return _jako_date(r[kd] if kd < len(r) else None, serial)
        try:
            return date(*(int(_jako_liczba(r[j])) for j in podzial))
        except (TypeError, ValueError, IndexError):
            return None

    # --- znaczniki czasu w kolumnie daty?
    znaczniki = [_jako_znacznik(r[kd]) for r in probka if kd is not None and kd < len(r)]
    ma_czas = sum(1 for t in znaczniki if t is not None and (t.hour, t.minute) != (0, 0)) >= 5

    # --- kolumna okresu: ≥ 90% komórek rozpoznanych, ≥ 20 kolejnych okresów w pierwszej pełnej dobie
    daty_probki, ostatnia = [], None
    for r in probka:
        d = doba(r) or ostatnia
        daty_probki.append(d)
        ostatnia = d
    grupy = Counter(daty_probki)
    pierwsza = next((d for d in dict.fromkeys(daty_probki) if d is not None and grupy[d] >= 20), None)
    wiersze_pierwszej = [r for r, d in zip(probka, daty_probki) if d == pierwsza]
    zajete = {kg, kd, *(podzial or ())}
    kandydaci = []
    for j in kolumny:
        if j in zajete:
            continue
        komorki = [r[j] for r in probka if j < len(r) and _tekst(r[j])]
        if len(komorki) < 20:
            continue
        rodzaj, ile = Counter((_okres(c) or ("?",))[0] for c in komorki).most_common(1)[0]
        if rodzaj == "?" or ile < 0.9 * len(komorki):
            continue
        wart = [o for o in (_okres_w_trybie(r[j] if j < len(r) else None, rodzaj) for r in wiersze_pierwszej) if o]
        rozne = sorted({o[1] for o in wart})
        if len(rozne) < 20:
            continue  # np. dzień tygodnia, numer miesiąca
        krok_k = _dominujacy_krok([b - a for a, b in zip(rozne, rozne[1:])]) or 1
        if not len(rozne) - 1 <= (rozne[-1] - rozne[0]) // krok_k + 1 <= len(rozne) + 2:
            continue  # wartości nie tworzą kolejnych okresów
        kandydaci.append((bool(RE_NAGL_OKRES.search(nagl[j])), -j, j, rodzaj))
    kandydaci.sort(reverse=True)
    ko, rodzaj_ko = (kandydaci[0][2], kandydaci[0][3]) if kandydaci else (None, None)
    tryb = rodzaj_ko if ko is not None else ("znacznik" if ma_czas else "kolejnosc")

    # --- wiersze danych
    rekordy = []  # (doba, znacznik, okres, wartość, pierwszy w grupie daty?)
    ostatnia = None
    for r in dane:
        if any(isinstance(c, str) and RE_POMIN.search(c) for c in r):
            continue
        surowa = r[kg] if kg < len(r) else None
        v = _jako_liczba(surowa)
        if v is None and _tekst(surowa):
            continue  # tekst w kolumnie G11 (powtórzony nagłówek, przypis)
        okres = None
        if ko is not None:
            komorka = r[ko] if ko < len(r) else None
            okres = _okres_w_trybie(komorka, rodzaj_ko)
            if okres is None and _tekst(komorka):
                continue  # nierozpoznana etykieta okresu
        d = doba(r)
        zn = _jako_znacznik(r[kd]) if tryb == "znacznik" and kd is not None and kd < len(r) else None
        nowa_doba = d is not None and d != ostatnia
        if d is None:
            if ostatnia is None or (v is None and okres is None):
                continue
            d = ostatnia  # scalone komórki daty
        if v is None and okres is None and zn is None and tryb != "kolejnosc":
            continue
        ostatnia = d
        rekordy.append((d, zn, okres, v, nowa_doba))
    if len(rekordy) < 20:
        return "za mało wierszy z wartościami G11"

    doby: dict[date, list[Wpis]] = defaultdict(list)
    if tryb == "znacznik":
        zn = [t for _, t, _, _, _ in rekordy if t is not None]
        krok = _dominujacy_krok([int((b - a).total_seconds() // 60) for a, b in zip(zn, zn[1:])])
        koniec = bool(zn and krok and zn[0].hour * 60 + zn[0].minute == krok)
        widziane: Counter = Counter()
        for _, t, _, v, _ in rekordy:
            if t is None:
                continue
            if koniec:
                n = widziane[t]
                widziane[t] += 1
                t = _poczatek_z_konca(t, n, krok)
            if t.year == rok:
                doby[t.date()].append(Wpis(t.hour * 60 + t.minute, None, v))
        opis = f"znacznik czasu w kolumnie {_litera(kd)}" + (" (koniec okresu)" if koniec else "")
        return opis, doby, "min", krok
    if tryb == "min":
        po_dniach: dict[date, list] = defaultdict(list)
        for d, _, o, v, _ in rekordy:
            po_dniach[d].append((o, v))
        roznice = []
        for lst in po_dniach.values():
            m = [o[1] for o, _ in lst if o]
            roznice += [b - a for a, b in zip(m, m[1:])]
        krok = _dominujacy_krok(roznice)
        pierwsza_lista = po_dniach[min(po_dniach)]
        koniec = bool(krok and pierwsza_lista and pierwsza_lista[0][0] and pierwsza_lista[0][0][1] == krok
                      and 0 not in {o[1] for o, _ in pierwsza_lista if o})
        for d, lst in po_dniach.items():
            widziane = Counter()
            for i, (o, v) in enumerate(lst):
                if o is None:
                    m, dd = None, d
                elif koniec:
                    # etykieta końca okresu -> lokalny koniec (00:00 na początku doby = koniec poprzedniej,
                    # 00:00 na końcu / 24:00 = koniec bieżącej) -> początek liczony w UTC (zmiana czasu)
                    lok = datetime(d.year, d.month, d.day) + timedelta(minutes=o[1])
                    if o[1] == 0 and i > 0:
                        lok += timedelta(days=1)
                    start = _poczatek_z_konca(lok, widziane[lok], krok)
                    widziane[lok] += 1
                    m, dd = start.hour * 60 + start.minute, start.date()
                else:
                    m, dd = o[1], d
                if dd.year == rok:
                    doby[dd].append(Wpis(m, None, v))
        opis = f"etykiety czasu w kolumnie {_litera(ko)}" + (" (koniec okresu)" if koniec else "")
        return opis, doby, "min", krok
    for d, _, o, v, _ in rekordy:
        if d.year == rok:
            doby[d].append(Wpis(None, (o[1], o[2]) if o else (len(doby[d]), 0), v))
    opis = f"numery okresów w kolumnie {_litera(ko)}" if tryb == "nr" else "kolejność wierszy"
    return opis, doby, "nr" if tryb == "nr" else "kolejnosc", _krok_z_liczby(doby)


# ------------------------------------------------------------------ układ szeroki

def _parsuj_szeroki(wiersze: list[tuple], rok: int, i: int):
    naglowek = wiersze[i]
    okresy = [(j, _okres(c)) for j, c in enumerate(naglowek) if not isinstance(c, (datetime, date))]
    okresy = [(j, o) for j, o in okresy if o is not None]
    if len(okresy) < 23:
        return None
    rodzaje = {o[0] for _, o in okresy}
    j0 = min(j for j, _ in okresy)
    dane = wiersze[i + 1:]
    probka = [r for r in dane[:400] if any(c is not None for c in r)]
    kd, serial = _kolumna_daty(probka, [_tekst(c) for c in naglowek], list(range(j0)))
    if kd is None:
        return "brak kolumny daty przed kolumnami okresów"
    kolejnosc = {j: ((o[1], o[2]) if len(rodzaje) == 1 else (n, 0)) for n, (j, o) in enumerate(okresy)}
    doby: dict[date, list[Wpis]] = {}
    for r in dane:
        d = _jako_date(r[kd] if kd < len(r) else None, serial)
        if d is None or d.year != rok or d in doby:
            continue
        doby[d] = [Wpis(None, kolejnosc[j], _jako_liczba(r[j]) if j < len(r) else None) for j, _ in okresy]
    opis = f"układ szeroki (nagłówek w wierszu {i + 1}, {len(okresy)} kolumn okresów)"
    return opis, doby, "nr", 15 if len(okresy) >= 88 else 60


# ------------------------------------------------------------------ cały plik

def _kandydaci(wiersze: list[tuple], rok: int, tytul: str):
    """Kolejne interpretacje arkusza: (opis, (dobre, złe)) albo (opis, powód odrzucenia)."""
    for i, w in enumerate(wiersze[:80]):
        teksty = [_tekst(c) for c in w]
        kol_g11 = [j for j, t in enumerate(teksty) if RE_GRUPA.search(t)]
        if not kol_g11:
            continue
        kg = next((j for j in kol_g11 if RE_GRUPA.fullmatch(teksty[j])), kol_g11[0])
        nad = wiersze[i - 1] if i > 0 else ()
        nagl = [" ".join(x for x in (_tekst(nad[j]) if j < len(nad) else "", teksty[j]) if x)
                for j in range(len(teksty))]
        prefiks = f"wiersz {i + 1}, G11 w kolumnie {_litera(kg)}"
        wynik = _parsuj_dlugi(wiersze[i + 1:], rok, nagl, kg)
        if isinstance(wynik, str):
            yield prefiks, wynik
            continue
        opis, doby, tryb, krok = wynik
        yield f"{prefiks}, {opis}", _ocen(doby, tryb, krok)
    arkusz_g11 = RE_GRUPA.search(tytul) or any(RE_GRUPA.search(_tekst(c)) for r in wiersze[:10] for c in r)
    if arkusz_g11:
        for i in range(min(40, len(wiersze))):
            wynik = _parsuj_szeroki(wiersze, rok, i)
            if wynik is None:
                continue
            if isinstance(wynik, str):
                yield f"wiersz {i + 1}", wynik
                continue
            opis, doby, tryb, krok = wynik
            yield opis, _ocen(doby, tryb, krok)
            break


def _akceptowalny(dobre: dict, zle: dict) -> bool:
    return len(dobre) >= 28 and len(zle) <= max(2, MAKS_WADLIWYCH * (len(dobre) + len(zle)))


def parsuj_profil(sciezka: Path, rok: int) -> list[list]:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(sciezka, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise BladProfilu(f"{sciezka.name} nie jest poprawnym plikiem XLSX ({e})") from e
    diagnoza: list[str] = []
    zebrane: dict[date, dict[int, float]] = {}
    zle_zebrane: dict[date, str] = {}
    arkusze = []
    for ws in wb.worksheets:
        wiersze = [tuple(r) for r in ws.iter_rows(values_only=True)]
        diagnoza.append(f"arkusz '{ws.title}': {len(wiersze)} wierszy; początek: {_podglad(wiersze)}")
        for opis, wynik in _kandydaci(wiersze, rok, ws.title):
            if isinstance(wynik, str):
                diagnoza.append(f"    {opis}: {wynik}")
                continue
            dobre, zle = wynik
            przyklady = "; ".join(f"{d}: {p}" for d, p in list(zle.items())[:3])
            diagnoza.append(f"    {opis}: {len(dobre)} dób poprawnych, {len(zle)} wadliwych"
                            + (f" (np. {przyklady})" if zle else ""))
            if not _akceptowalny(dobre, zle):
                continue
            if len(dobre) >= 300:
                return _wynik_roku(dobre, zle, rok, f"Enea {rok} (arkusz {ws.title}; {opis})")
            for d, s in dobre.items():
                zebrane.setdefault(d, s)
            for d, p in zle.items():
                zle_zebrane.setdefault(d, p)
            arkusze.append(ws.title)
            break
    if len(zebrane) >= 300:
        zle_zebrane = {d: p for d, p in zle_zebrane.items() if d not in zebrane}
        return _wynik_roku(zebrane, zle_zebrane, rok, f"Enea {rok} (arkusze: {', '.join(arkusze)})")
    raise BladProfilu("Nie rozpoznano układu pliku profilu. Diagnoza:\n  " + "\n  ".join(diagnoza))


def _wynik_roku(dobre: dict[date, dict[int, float]], zle: dict[date, str], rok: int, zrodlo: str) -> list[list]:
    dni = list(zakres_dni(date(rok, 1, 1), date(rok, 12, 31)))
    brak = [d for d in dni if d not in dobre]
    if len(brak) > MAKS_WADLIWYCH * len(dni):
        raise BladProfilu(f"profil {rok}: brak poprawnych danych dla {len(brak)} dób ({zrodlo})")
    profil = {d: dict(dobre[d]) for d in dni if d in dobre}
    if brak:
        wzor: dict[tuple, list[float]] = defaultdict(list)
        for d, sloty in profil.items():
            for s, v in sloty.items():
                wzor[(d.month, typ_dnia(d), s)].append(v)
        for d in brak:
            profil[d] = {s: statistics.fmean(wzor.get((d.month, typ_dnia(d), s))
                                             or wzor.get((d.month, "roboczy", s)) or [0.0])
                         for s in range(1, 97)}
        log.warning("OSD: profil %d – %d dób uzupełniono średnią (miesiąc, typ dnia, slot): %s", rok, len(brak),
                    ", ".join(f"{d}{' – ' + zle[d] if d in zle else ''}" for d in brak[:8]))
    razem = sum(v for sloty in profil.values() for v in sloty.values())
    if razem <= 0:
        raise BladProfilu(f"profil {rok}: suma roku = 0")
    OPISY[rok] = zrodlo
    tag = f"Enea {rok}, parser v{WERSJA_PARSERA}"  # krótko – kolumna powtarza się w każdym wierszu
    return [[d.isoformat(), s, fmt_liczba(profil[d][s] * SUMA_ROKU / razem, 6),
             tag if d in dobre else f"{tag}; doba uzupełniona"]
            for d in dni for s in sorted(profil[d])]


def opis_ksztaltu(wiersze: list[list]) -> str:
    """Godzina szczytu i minimum średniego profilu – szybka kontrola sensowności (G11: szczyt wieczorem)."""
    godziny: dict[int, float] = defaultdict(float)
    for _, slot, v, _ in wiersze:
        godziny[(int(slot) - 1) // 4] += float(v)
    razem = sum(godziny.values()) or 1
    szczyt = max(godziny, key=godziny.get)
    dolek = min(godziny, key=godziny.get)
    return (f"szczyt {szczyt:02d}:00–{szczyt + 1:02d}:00 ({100 * godziny[szczyt] / razem:.1f}% doby), "
            f"minimum {dolek:02d}:00–{dolek + 1:02d}:00 ({100 * godziny[dolek] / razem:.1f}%)")


# ------------------------------------------------------------------ uzupełnianie brakujących lat

def przyblizony_profil(rok: int, wzor: list[list], rok_wzoru: int) -> list[list]:
    """Profil roku bez pliku OSD: średnia z roku-wzoru wg (miesiąc, typ dnia, slot)."""
    sumy: dict[tuple, list[float]] = defaultdict(list)
    for d_s, slot, v, _ in wzor:
        d = date.fromisoformat(d_s)
        sumy[(d.month, typ_dnia(d), int(slot))].append(float(v))
    srednie = {k: sum(v) / len(v) for k, v in sumy.items()}
    wartosci = {}
    for d in zakres_dni(date(rok, 1, 1), date(rok, 12, 31)):
        for slot in range(1, 97):
            v = srednie.get((d.month, typ_dnia(d), slot))
            if v is None:  # np. brak takiego typu dnia w danym miesiącu wzoru
                v = srednie.get((d.month, "roboczy", slot), 0.0)
            wartosci[(d, slot)] = v
    razem = sum(wartosci.values()) or 1
    return [[d.isoformat(), slot, fmt_liczba(v * SUMA_ROKU / razem, 6), f"przybliżenie z profilu {rok_wzoru}"]
            for (d, slot), v in wartosci.items()]


def _aktualny(cel: Path) -> bool:
    wiersze = czytaj_csv(cel)[:5]
    return bool(wiersze) and all(w["zrodlo"].startswith("Enea") and f"parser v{WERSJA_PARSERA}" in w["zrodlo"]
                                 for w in wiersze)


def aktualizuj(lata: list[int], wymus: bool = False) -> dict:
    sesja = None
    gotowe: dict[int, Path] = {}
    bledy: dict[int, str] = {}
    for rok in sorted(set(lata) | set(PROFILE_URL)):
        if rok not in PROFILE_URL:
            continue
        cel = plik_profilu(rok)
        if not wymus and _aktualny(cel):
            gotowe[rok] = cel
            continue
        surowy = plik_surowy(rok)
        try:
            if not surowy.exists() or wymus:
                sesja = sesja or nowa_sesja()
                log.info("OSD: pobieram profil Enea %d", rok)
                r = pobierz(sesja, PROFILE_URL[rok], timeout=120)
                if not r.content.startswith(b"PK"):
                    raise BladProfilu(f"{PROFILE_URL[rok]} nie zwrócił pliku XLSX "
                                      f"(Content-Type: {r.headers.get('Content-Type')}) – sprawdź adres na stronie IRiESD")
                surowy.parent.mkdir(parents=True, exist_ok=True)
                surowy.write_bytes(r.content)
            wiersze = parsuj_profil(surowy, rok)
        except (BladProfilu, BladPobierania, requests.RequestException) as e:
            log.error("OSD: profil %d – %s", rok, e)
            bledy[rok] = str(e).splitlines()[0]
            continue
        zapisz_csv_atomowo(cel, NAGLOWEK, wiersze)
        log.info("OSD: profil %d – %d wierszy; %s; %s", rok, len(wiersze), opis_ksztaltu(wiersze), OPISY.get(rok))
        zapisz_stan({f"profil_osd_{rok}": OPISY.get(rok, "")})
        gotowe[rok] = cel
    # lata z danymi rynkowymi, dla których OSD nie opublikował (jeszcze) profilu albo plik nie dał się odczytać
    for rok in sorted((set(lata) | set(bledy)) - set(gotowe)):
        rok_wzoru = min(gotowe, key=lambda r: abs(r - rok)) if gotowe else None
        if rok_wzoru is None:
            raise BladProfilu("brak jakiegokolwiek poprawnego profilu OSD do przybliżenia"
                              + (f" (błędy: {bledy})" if bledy else ""))
        wzor = [[w["data"], w["slot"], w["wartosc"], w["zrodlo"]] for w in czytaj_csv(gotowe[rok_wzoru])]
        zapisz_csv_atomowo(plik_profilu(rok), NAGLOWEK, przyblizony_profil(rok, wzor, rok_wzoru))
        log.warning("OSD: brak profilu Enea na %d – zapisano przybliżenie z roku %d", rok, rok_wzoru)
    if bledy:
        raise BladProfilu(f"nie odczytano profilu Enea dla lat {sorted(bledy)} – użyto przybliżenia; "
                          "szczegóły wyżej w logu")
    return {"lata": sorted(set(lata) | set(gotowe))}
