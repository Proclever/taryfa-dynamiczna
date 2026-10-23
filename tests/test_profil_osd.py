"""Parser profilu OSD na wielu możliwych układach arkusza (pliki XLSX generowane w teście, bez sieci).

Wartości syntetyczne mają szczyt o 19:00 i dołek w nocy, więc błędne przypisanie godzin
(przesunięcie, cała doba w jednym slocie) od razu wychodzi w asercjach.
"""
from __future__ import annotations

import math
import sys
from datetime import date, datetime, time, timedelta
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))

import zrodla_osd  # noqa: E402
from wspolne import TZ_PL, poczatek_doby_utc  # noqa: E402

ROK = 2025
WIOSNA, JESIEN = date(2025, 3, 30), date(2025, 10, 26)


def dni():
    d = date(ROK, 1, 1)
    while d.year == ROK:
        yield d
        d += timedelta(days=1)


def ksztalt(godzina: float) -> float:
    return 1e-4 * (1 + 0.9 * math.exp(-((godzina - 19.25) / 2.2) ** 2) + 0.3 * math.exp(-((godzina - 8) / 1.5) ** 2))


def okresy(d: date, krok: int = 60):
    """Okresy doby w kolejności UTC: (lokalny początek, lokalny koniec, wartość)."""
    t, koniec = poczatek_doby_utc(d), poczatek_doby_utc(d + timedelta(days=1))
    wynik = []
    while t < koniec:
        a, b = t.astimezone(TZ_PL), (t + timedelta(minutes=krok)).astimezone(TZ_PL)
        wynik.append((a.replace(tzinfo=None), b.replace(tzinfo=None), ksztalt(a.hour + a.minute / 60) * krok / 60))
        t += timedelta(minutes=krok)
    return wynik


def dt(d: date) -> datetime:
    return datetime(d.year, d.month, d.day)


# ------------------------------------------------------------------ układy

def u_numer(wb, ws, tytul="Załącznik nr 4 – standardowe profile zużycia", suma=False, scalone=False, pusta=None):
    ws.append([tytul])
    ws.append([])
    ws.append(["Data", "Godzina", "G11", "G12", "G12w"])
    for d in dni():
        lst = okresy(d)
        for i, (a, _, v) in enumerate(lst):
            if pusta == (d, i):
                v = None
            ws.append([dt(d) if (i == 0 or not scalone) else None, i + 1, v, 0.0002, 0.0003])
        if suma:
            ws.append([None, "Suma", sum(v for _, _, v in lst), None, None])


def u_2a(wb, ws):
    ws.append(["Data", "Godz.", "G11"])
    for d in dni():
        widziane = set()
        for a, _, v in okresy(d):
            etykieta = f"{a.hour}A" if a.hour in widziane else str(a.hour)
            widziane.add(a.hour)
            ws.append([d.strftime("%d.%m.%Y"), etykieta, v])


def u_dzien_tygodnia(wb, ws):
    ws.append(["Dzień tygodnia", "Data", "Godzina", "G11", "G12"])
    for d in dni():
        for i, (_, _, v) in enumerate(okresy(d)):
            ws.append([["pn", "wt", "śr", "cz", "pt", "so", "nd"][d.weekday()], dt(d), i + 1, v, 0.0002])


def u_rok_mies_dzien(wb, ws):
    ws.append(["Rok", "Miesiąc", "Dzień", "Godzina", "G11"])
    for d in dni():
        for i, (_, _, v) in enumerate(okresy(d)):
            ws.append([d.year, d.month, d.day, i + 1, v])


def u_etykiety_konca(wb, ws):
    ws.append(["Data", "Godzina", "G11"])
    for d in dni():
        for _, b, v in okresy(d):
            ws.append([dt(d), "24:00" if b.date() > d else f"{b.hour:02d}:00", v])


def u_zakresy(wb, ws):
    ws.append(["Data", "Okres", "G11"])
    for d in dni():
        for a, b, v in okresy(d):
            ws.append([dt(d), f"{a.hour:02d}:00-{(a.hour + 1) % 24:02d}:00", v])


def u_miesiace(wb, ws):
    wb.remove(ws)
    for m in range(1, 13):
        s = wb.create_sheet(f"{m:02d}.{ROK}")
        s.append(["Data", "Godzina", "G11", "G12"])
        for d in dni():
            if d.month == m:
                for i, (_, _, v) in enumerate(okresy(d)):
                    s.append([dt(d), i + 1, v, 0.0002])


def u_szeroki_25(wb, ws):
    ws.title = "G11"
    ws.append(["Data"] + list(range(1, 26)))
    for d in dni():
        wart = [v for _, _, v in okresy(d)]
        if len(wart) == 23:  # wiosna: godzina 2:00–3:00 nie istnieje
            wart = wart[:2] + [None] + wart[2:]
        ws.append([dt(d)] + wart + [None] * (25 - len(wart)))


def u_szeroki_2a(wb, ws):
    ws.title = "Profil G11"
    ws.append(["Data"] + [str(h) for h in range(24)] + ["2A"])
    for d in dni():
        lst = okresy(d)
        po_godz = {}
        extra = None
        for a, _, v in lst:
            if a.hour in po_godz:
                extra = v
            else:
                po_godz[a.hour] = v
        ws.append([dt(d)] + [po_godz.get(h) for h in range(24)] + [extra])


def u_znaczniki_konca(wb, ws):
    ws.append(["Data i godzina", "G11"])
    for d in dni():
        for _, b, v in okresy(d):
            ws.append([b, v])


def u_znaczniki_15(wb, ws):
    ws.append(["Data", "G11", "G12"])
    for d in dni():
        for a, _, v in okresy(d, 15):
            ws.append([a, v, 0.0001])


def u_dwuwierszowy(wb, ws):
    ws.append(["Data", "Godzina", "Grupa taryfowa", None, None])
    ws.append([None, None, "G11", "G12", "G12w"])
    ws.merge_cells("A1:A2")
    ws.merge_cells("B1:B2")
    ws.merge_cells("C1:E1")
    for d in dni():
        for i, (_, _, v) in enumerate(okresy(d)):
            ws.append([dt(d), i + 1, v, 0.0002, 0.0003])


def u_kwadranse_czas(wb, ws):
    ws.append(["Data", "Czas", "G11"])
    for d in dni():
        for a, _, v in okresy(d, 15):
            ws.append([dt(d), time(a.hour, a.minute), v])


def u_daty_seryjne(wb, ws):
    ws.append(["Data", "Godzina", "G11"])
    for d in dni():
        for i, (_, _, v) in enumerate(okresy(d)):
            ws.append([(d - date(1899, 12, 30)).days, i + 1, v])


def u_enea(wb, ws):
    """Układ plików Enea (IRiESD, zał. 4): arkusz na grupę, wiersze = doby, kolumny 1–24 + „2a”.

    Godziny numerowane wg końca okresu (odczyt zegara): wiosną kolumna „2” = 0, a 1:00–2:00 CET
    jest w kolumnie „3”; jesienią „2a” to 2:00–3:00 czasu letniego (między „2” a „3”).
    """
    wb.remove(ws)
    for tytul in ("B11 B11em B12 B12eco B13active", "G11 G11pewna", "G12 G12eco G13active "):
        s = wb.create_sheet(tytul)
        s.append(["Dzień", "Data"] + list(range(1, 25)) + ["2a", None])
        for d in dni():
            lst = [v for _, _, v in okresy(d)]
            if len(lst) == 23:
                kol, extra = [lst[0], 0] + lst[1:], None
            elif len(lst) == 25:
                kol, extra = lst[:2] + lst[3:], lst[2]
            else:
                kol, extra = lst, None
            mnoznik = 1 if tytul.startswith("G11") else 0.5  # inne grupy mają inny kształt
            s.append([["pn", "wt", "śr", "cz", "pt", "so", "nd"][d.weekday()], dt(d)]
                     + [v * mnoznik if v else v for v in kol] + [extra, None])
        s.append([None] * 28)
        s.append([None] * 28)


UKLADY = {
    "enea": (u_enea, True),
    "numer": (u_numer, True),
    "tytul_z_G11": (lambda wb, ws: u_numer(wb, ws, tytul="Profile G11, G12, G12w na rok 2025"), True),
    "suma_dobowa": (lambda wb, ws: u_numer(wb, ws, suma=True), True),
    "scalone_daty": (lambda wb, ws: u_numer(wb, ws, scalone=True), True),
    "pusta_komorka": (lambda wb, ws: u_numer(wb, ws, pusta=(date(2025, 2, 12), 10)), True),
    "etykiety_2A": (u_2a, True),
    "dzien_tygodnia": (u_dzien_tygodnia, True),
    "rok_mies_dzien": (u_rok_mies_dzien, True),
    "etykiety_konca": (u_etykiety_konca, True),
    "zakresy": (u_zakresy, True),
    "arkusz_na_miesiac": (u_miesiace, True),
    "szeroki_25": (u_szeroki_25, True),
    "szeroki_2A": (u_szeroki_2a, True),
    "znaczniki_konca": (u_znaczniki_konca, True),
    "znaczniki_15min": (u_znaczniki_15, True),
    "naglowek_dwuwierszowy": (u_dwuwierszowy, True),
    "kwadranse_czas": (u_kwadranse_czas, True),
    "daty_seryjne": (u_daty_seryjne, True),
}


@pytest.mark.parametrize("nazwa", list(UKLADY))
def test_uklady_profilu(tmp_path, nazwa):
    from openpyxl import Workbook
    budowa, jesien_dokladna = UKLADY[nazwa]
    wb = Workbook()
    ws = wb.active
    ws.title = "Profile"
    budowa(wb, ws)
    p = tmp_path / f"{nazwa}.xlsx"
    wb.save(p)

    wiersze = zrodla_osd.parsuj_profil(p, ROK)
    po_dniach: dict[str, dict[int, float]] = {}
    for d, s, v, _ in wiersze:
        po_dniach.setdefault(d, {})[int(s)] = float(v)
    assert len(po_dniach) == 365
    assert abs(sum(sum(x.values()) for x in po_dniach.values()) - zrodla_osd.SUMA_ROKU) < 1
    uzupelnione = {w[0] for w in wiersze if "uzupełniona" in w[3]}
    assert len(uzupelnione) <= 2, uzupelnione

    for dzien in ("2025-01-15", "2025-07-15"):
        godz = {}
        for s, v in po_dniach[dzien].items():
            godz[(s - 1) // 4] = godz.get((s - 1) // 4, 0) + v
        assert max(godz, key=godz.get) == 19, (dzien, godz)       # szczyt wieczorny na swoim miejscu
        assert min(godz, key=godz.get) in range(0, 6), godz       # dołek w nocy
        assert len(godz) == 24

    wiosna = po_dniach[WIOSNA.isoformat()]
    assert not any(wiosna.get(s, 0) > 0 for s in range(9, 13))       # 02:00–02:59 nie istnieje
    # godzina 1:00–2:00 ma swoją energię (nie zero, nie przesunięta na nieistniejącą 2:00)
    assert wiosna[5] / wiosna[13] == pytest.approx(ksztalt(1) / ksztalt(3), rel=1e-3)
    if jesien_dokladna:
        jesien = po_dniach[JESIEN.isoformat()]
        assert JESIEN.isoformat() not in uzupelnione
        # godzina 2:00–3:00 występuje dwa razy -> podwójna energia w slotach 9–12
        assert jesien[9] / jesien[13] == pytest.approx(2 * ksztalt(2) / ksztalt(3), rel=1e-3)


def test_nierozpoznany_uklad_daje_diagnoze(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Opis"
    ws.append(["To nie jest profil"])
    ws.append(["G11", "tylko", "tekst"])
    p = tmp_path / "zly.xlsx"
    wb.save(p)
    with pytest.raises(zrodla_osd.BladProfilu, match="Diagnoza"):
        zrodla_osd.parsuj_profil(p, ROK)


def test_cala_doba_w_jednym_slocie_jest_odrzucana(tmp_path):
    """Kolumna godziny bez rozpoznawalnych wartości: nie wolno wpisać całej doby do 00:00."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Data", "Uwagi", "G11"])
    for d in dni():
        for _, _, v in okresy(d):
            ws.append([dt(d), "x", v])
    p = tmp_path / "bez_godzin.xlsx"
    wb.save(p)
    wiersze = zrodla_osd.parsuj_profil(p, ROK)  # kolejność wierszy = kolejne godziny
    godz = {}
    for d, s, v, _ in wiersze:
        if d == "2025-01-15":
            godz[(int(s) - 1) // 4] = godz.get((int(s) - 1) // 4, 0) + float(v)
    assert max(godz, key=godz.get) == 19


def test_ksztalt_w_logu():
    wiersze = [["2025-01-01", s, "1" if s != 77 else "5", "x"] for s in range(1, 97)]
    assert zrodla_osd.opis_ksztaltu(wiersze).startswith("szczyt 19:00–20:00")
